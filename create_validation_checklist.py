"""
Generate CloudGuard Validation Checklist for Stakeholder Review.
Creates a comprehensive Excel workbook with feature validation items
organized by stakeholder domain (VP Networking, Cloud, Application, CISO).
"""

import openpyxl
from openpyxl.styles import (
    Font, PatternFill, Alignment, Border, Side, NamedStyle
)
from openpyxl.utils import get_column_letter
from datetime import datetime


def create_checklist():
    wb = openpyxl.Workbook()

    # ── Colour palette ──
    DARK_BLUE = "1B2A4A"
    MEDIUM_BLUE = "2D4A7A"
    LIGHT_BLUE = "D6E4F0"
    WHITE = "FFFFFF"
    LIGHT_GRAY = "F2F2F2"
    GREEN = "27AE60"
    AMBER = "F39C12"
    RED = "E74C3C"
    HEADER_BG = PatternFill("solid", fgColor=DARK_BLUE)
    SECTION_BG = PatternFill("solid", fgColor=MEDIUM_BLUE)
    ALT_ROW = PatternFill("solid", fgColor=LIGHT_GRAY)
    PASS_FILL = PatternFill("solid", fgColor=GREEN)
    FAIL_FILL = PatternFill("solid", fgColor=RED)

    HEADER_FONT = Font(name="Calibri", size=11, bold=True, color=WHITE)
    SECTION_FONT = Font(name="Calibri", size=11, bold=True, color=WHITE)
    BODY_FONT = Font(name="Calibri", size=10, color="333333")
    TITLE_FONT = Font(name="Calibri", size=16, bold=True, color=DARK_BLUE)
    SUBTITLE_FONT = Font(name="Calibri", size=11, color="666666")

    thin_border = Border(
        left=Side(style="thin", color="D0D0D0"),
        right=Side(style="thin", color="D0D0D0"),
        top=Side(style="thin", color="D0D0D0"),
        bottom=Side(style="thin", color="D0D0D0"),
    )

    wrap_align = Alignment(wrap_text=True, vertical="center")
    center_align = Alignment(horizontal="center", vertical="center", wrap_text=True)

    # ═══════════════════════════════════════════
    #  COLUMNS: S.No | Feature | Description | Stakeholder | Status | Validated By | Date | Remarks
    # ═══════════════════════════════════════════
    HEADERS = [
        "S.No",
        "Feature / Capability",
        "Description & Validation Criteria",
        "Relevant Stakeholder",
        "Status\n(Pass / Fail / N/A)",
        "Validated By",
        "Validation Date",
        "Remarks / Observations",
    ]
    COL_WIDTHS = [7, 32, 55, 22, 18, 18, 16, 35]

    # ── Feature data per stakeholder section ──
    # (section_title, [(feature, description, stakeholder), ...])

    SECTIONS = [
        # ─────────────── VP NETWORKING ───────────────
        (
            "A. Network & Infrastructure Monitoring (VP – Networking)",
            [
                (
                    "Azure Front Door Log Ingestion",
                    "Verify that CloudGuard pulls access and WAF logs from Azure Front Door via Azure Monitor REST API. "
                    "Confirm logs include client IP, HTTP status code, URL path, routing rule, and time-taken-ms.",
                    "VP – Networking",
                ),
                (
                    "Azure Application Gateway Log Ingestion",
                    "Verify that L7 Application Gateway access logs (HTTP method, URL, response code, backend latency) are "
                    "collected and stored in the CloudGuard database.",
                    "VP – Networking",
                ),
                (
                    "Azure API Management (APIM) Log Ingestion",
                    "Confirm APIM gateway logs are ingested with fields such as API operation, response code, backend time, "
                    "and subscription key identifier.",
                    "VP – Networking",
                ),
                (
                    "Azure VM Diagnostic Log Ingestion",
                    "Verify that Azure VM diagnostic / activity logs are collected from the configured VM resource ID.",
                    "VP – Networking",
                ),
                (
                    "Real-Time Service Health Checks",
                    "Validate that the Health Check module pings Front Door, App Gateway, APIM, App Service, Blob Storage, "
                    "WAF, and Log Analytics endpoints in real time and shows HTTP status code, response time (ms), "
                    "and healthy/unhealthy/unknown status on the dashboard.",
                    "VP – Networking",
                ),
                (
                    "Infrastructure Flow Diagram",
                    "Confirm the dashboard displays an interactive infrastructure topology diagram showing "
                    "the flow: Front Door → App Gateway → APIM → App Service with live health indicators.",
                    "VP – Networking",
                ),
                (
                    "Network Latency & Response Time Tracking",
                    "Verify that log preprocessing extracts and normalises 'time_taken' and backend latency fields, "
                    "and that these are surfaced in analytics charts (e.g., Golden Signals – Latency).",
                    "VP – Networking",
                ),
            ],
        ),
        # ─────────────── VP CLOUD ───────────────
        (
            "B. Cloud Operations & Platform (VP – Cloud)",
            [
                (
                    "Azure Monitor REST API Integration",
                    "Validate that the Log Collector agent authenticates via Azure AD Service Principal "
                    "(Tenant ID, Client ID, Client Secret) and pulls activity logs from Azure Monitor for all "
                    "configured resource IDs (Front Door, App Gateway, APIM, VM).",
                    "VP – Cloud",
                ),
                (
                    "Azure PostgreSQL Database Connectivity",
                    "Confirm the application connects to Azure Database for PostgreSQL (async via asyncpg) and "
                    "that all tables (RawLog, CloudLog, Incident, AgentRun, RootCauseAnalysis, IncidentGroup) "
                    "are created successfully on startup.",
                    "VP – Cloud",
                ),
                (
                    "Automated Pipeline Scheduler",
                    "Verify the APScheduler-based scheduler runs the 9-agent pipeline automatically every N hours "
                    "(configurable via SCHEDULER_INTERVAL_HOURS). Confirm enable/disable toggle works.",
                    "VP – Cloud",
                ),
                (
                    "Manual Pipeline Trigger",
                    "Validate that the 'Run Pipeline' button on the dashboard triggers the full 9-agent pipeline "
                    "on demand and returns results.",
                    "VP – Cloud",
                ),
                (
                    "Pipeline Enable / Disable Toggle",
                    "Confirm the pipeline master switch (PIPELINE_ENABLED) can be toggled from the dashboard UI "
                    "and the Settings page, and that a disabled pipeline skips scheduled runs.",
                    "VP – Cloud",
                ),
                (
                    "Docker Compose Deployment",
                    "Verify the application deploys successfully on an Azure VM using 'docker compose up -d' "
                    "with three containers: backend (FastAPI), frontend (React/Nginx), and Nginx reverse proxy.",
                    "VP – Cloud",
                ),
                (
                    "Google Cloud Vertex AI (Gemini) Integration",
                    "Confirm the Classification, RCA, Resolution, and Notification agents call Gemini 2.5 Flash "
                    "via Vertex AI using the configured GCP Service Account Key. Validate API responses.",
                    "VP – Cloud",
                ),
                (
                    "Pipeline Run History & Audit Trail",
                    "Verify that every pipeline run (manual or scheduled) is logged in the AgentRun table with "
                    "run ID, trigger type, start time, duration, logs processed, incidents created, P1/P2/P3 counts, "
                    "emails sent, and status (SUCCESS / FAILED).",
                    "VP – Cloud",
                ),
                (
                    "Log Analytics Workspace Integration",
                    "Validate that the Azure Log Analytics Workspace ID is used for health checks "
                    "and that the endpoint responds correctly.",
                    "VP – Cloud",
                ),
            ],
        ),
        # ─────────────── VP APPLICATION ───────────────
        (
            "C. Application & AI Pipeline (VP – Application)",
            [
                (
                    "9-Agent LangGraph Pipeline Execution",
                    "Verify the full pipeline executes all 9 agents in order: Log Collector → Preprocessing → "
                    "Classification → RCA → Priority → Context → Resolution → Orchestrator → Notification. "
                    "Confirm each agent produces expected output.",
                    "VP – Application",
                ),
                (
                    "Log Preprocessing & Deduplication",
                    "Validate that the Preprocessing Engine cleans raw logs, normalises timestamps, extracts "
                    "structured fields (level, category, source, resource ID), and removes duplicate entries.",
                    "VP – Application",
                ),
                (
                    "AI-Powered Incident Classification (Gemini)",
                    "Confirm the Classification Agent uses Gemini AI to detect incident type, intent, severity, "
                    "and categorisation from preprocessed log data.",
                    "VP – Application",
                ),
                (
                    "Root Cause Analysis (RCA) Agent",
                    "Verify the RCA agent identifies probable root causes with confidence scores, collects evidence, "
                    "matches historical incidents, and identifies affected components "
                    "(Backend Timeout, WAF Block, SSL Error, etc.).",
                    "VP – Application",
                ),
                (
                    "Priority Assignment (P1 / P2 / P3)",
                    "Validate that the Priority Agent assigns correct priority levels: "
                    "P1 (Critical – immediate action), P2 (Warning – monitor), P3 (Informational – log only). "
                    "Verify priority justification is recorded.",
                    "VP – Application",
                ),
                (
                    "Historical Context Lookup",
                    "Confirm the Context Agent queries the PostgreSQL database for similar past incidents "
                    "and surfaces historical resolution patterns.",
                    "VP – Application",
                ),
                (
                    "AI-Powered Resolution Recommendations",
                    "Verify the Resolution Agent generates actionable fix recommendations and runbooks "
                    "using Gemini AI, including step-by-step remediation instructions.",
                    "VP – Application",
                ),
                (
                    "Orchestrator Summary & Routing",
                    "Confirm the Orchestrator Agent aggregates outputs from all previous agents, "
                    "builds a pipeline summary, and routes notifications based on priority.",
                    "VP – Application",
                ),
                (
                    "Incident Detail Panel (UI)",
                    "Validate the incident detail panel shows: classification, priority justification, "
                    "RCA with confidence score, historical context, AI-generated resolution, and runbook.",
                    "VP – Application",
                ),
                (
                    "Real-Time WebSocket Updates",
                    "Confirm the dashboard receives live pipeline progress updates via WebSocket: "
                    "node_active, node_complete, pipeline_start, pipeline_complete, pipeline_error events.",
                    "VP – Application",
                ),
                (
                    "Workflow Visualization (React Flow)",
                    "Verify the Workflow View page displays the 9-agent pipeline as an interactive graph "
                    "with animated node status transitions (idle → running → completed), "
                    "real-time output result cards, and a pipeline summary modal on completion.",
                    "VP – Application",
                ),
                (
                    "Dashboard KPI Cards",
                    "Validate top-level KPI cards show: Total Incidents, Open Incidents, P1/P2/P3 counts, "
                    "Resolved count, Avg Resolution Time, Emails Sent, and pipeline run status.",
                    "VP – Application",
                ),
                (
                    "Analytics & Charting",
                    "Verify the Analytics page displays: Incident Trend (line chart over N days), "
                    "Service Breakdown (by Azure service), Error Distribution (by HTTP status/log level), "
                    "Notification History, Top Recurring Issues, MTTR metrics, Golden Signals, "
                    "Root Cause Distribution, Component Breakdown, and Resolution Metrics.",
                    "VP – Application",
                ),
                (
                    "Log Explorer with Search & Filters",
                    "Confirm the Log Explorer page provides paginated log viewing with filters "
                    "for log level, source service, and free-text search.",
                    "VP – Application",
                ),
                (
                    "Data Export (JSON / CSV / Excel)",
                    "Validate that logs and incidents can be exported in JSON, CSV, and XLSX formats "
                    "via the Export API and download buttons on the Log Explorer page.",
                    "VP – Application",
                ),
                (
                    "Incident Management (CRUD)",
                    "Confirm incidents can be listed with pagination, filtered by priority/status/category/service, "
                    "viewed in detail, and updated (e.g., mark as Resolved).",
                    "VP – Application",
                ),
            ],
        ),
        # ─────────────── CISO ───────────────
        (
            "D. Security & Compliance (CISO)",
            [
                (
                    "JWT Authentication & Protected API",
                    "Verify that all API endpoints (except /api/auth/login) require a valid JWT Bearer token. "
                    "Confirm login with email/password returns a signed JWT with 24-hour expiry.",
                    "CISO",
                ),
                (
                    "Login / Logout Flow (UI)",
                    "Validate the login page authenticates the user and redirects to the dashboard. "
                    "Confirm logout clears the JWT token and redirects to login.",
                    "CISO",
                ),
                (
                    "Protected Route Enforcement",
                    "Confirm that unauthenticated users cannot access the Dashboard, Analytics, "
                    "Workflow, Log Explorer, or Settings pages — they are redirected to Login.",
                    "CISO",
                ),
                (
                    "WAF Log Monitoring (Azure Front Door WAF)",
                    "Verify that WAF logs are ingested and that WAF-blocked requests (403 status, "
                    "rule violations) are classified as security incidents with appropriate priority.",
                    "CISO",
                ),
                (
                    "SSL/TLS Certificate Error Detection",
                    "Confirm the RCA agent detects SSL/TLS certificate errors (expired, mismatch, "
                    "handshake failure) from log patterns and raises them as security-relevant incidents.",
                    "CISO",
                ),
                (
                    "P1 Critical Incident Email Alerting",
                    "Validate that P1 (Critical) incidents trigger immediate email notifications "
                    "via the configured channel (SMTP / Azure Communication Service / Microsoft Graph API). "
                    "Confirm email contains incident title, priority, service, description, and recommended fix.",
                    "CISO",
                ),
                (
                    "Priority-Based Notification Routing",
                    "Confirm notification routing follows the policy: "
                    "P1 → Immediate email, P2 → Email digest, P3 → Dashboard only (no email).",
                    "CISO",
                ),
                (
                    "Email Notification Audit (DB Records)",
                    "Verify that every email sent is recorded in the database with: "
                    "email_sent flag, email_sent_at timestamp, and email_recipient address.",
                    "CISO",
                ),
                (
                    "CORS Policy Configuration",
                    "Confirm CORS origins are restricted to configured allowed origins only "
                    "(CORS_ORIGINS setting) and are not set to wildcard '*' in production.",
                    "CISO",
                ),
                (
                    "Sensitive Credential Management",
                    "Verify that all sensitive credentials (Azure Client Secret, JWT Secret Key, "
                    "SMTP Password, GCP Service Account Key) are stored in .env file and not "
                    "hardcoded in source code. Confirm .gitignore excludes .env and key files.",
                    "CISO",
                ),
            ],
        ),
        # ─────────────── CROSS-FUNCTIONAL ───────────────
        (
            "E. Cross-Functional / Overall System Validation",
            [
                (
                    "End-to-End Pipeline Smoke Test",
                    "Run the full pipeline end-to-end (manual trigger) and verify: logs are collected, "
                    "preprocessed, classified, analysed for root cause, prioritised, enriched with context, "
                    "resolutions generated, summary orchestrated, and notifications dispatched.",
                    "All Stakeholders",
                ),
                (
                    "System Health API",
                    "Verify the root endpoint (GET /) returns application name, version, and 'healthy' status.",
                    "All Stakeholders",
                ),
                (
                    "Dashboard Summary API",
                    "Confirm GET /api/dashboard returns combined summary: raw log stats, processed log stats, "
                    "incident counts (total, open, resolved, P1/P2/P3), emails sent, avg resolution time, "
                    "service-wise breakdown, and pipeline run status.",
                    "All Stakeholders",
                ),
                (
                    "Multi-Notification Channel Support",
                    "Verify that the notification system supports three channels: "
                    "SMTP, Azure Communication Service, and Microsoft Graph API — "
                    "configurable via NOTIFICATION_CHANNEL setting.",
                    "All Stakeholders",
                ),
                (
                    "Responsive & Modern UI Design",
                    "Confirm the React dashboard is responsive, visually polished, and includes: "
                    "sidebar navigation, header bar, dark-themed charts, animated counters, "
                    "sparkline graphs, and smooth page transitions.",
                    "All Stakeholders",
                ),
                (
                    "Settings Page — Agent Status & History",
                    "Verify the Settings page shows pipeline toggle, current agent status, "
                    "and a table of recent pipeline run history with trigger type, logs processed, "
                    "P1/P2/P3 counts, duration, and status.",
                    "All Stakeholders",
                ),
            ],
        ),
    ]

    # ═══════════════════════════════════════════
    #  BUILD THE WORKSHEET
    # ═══════════════════════════════════════════
    ws = wb.active
    ws.title = "Validation Checklist"

    # ── Title rows ──
    ws.merge_cells("A1:H1")
    title_cell = ws["A1"]
    title_cell.value = "CloudGuard — Stakeholder Validation Checklist"
    title_cell.font = TITLE_FONT
    title_cell.alignment = Alignment(horizontal="center", vertical="center")
    ws.row_dimensions[1].height = 36

    ws.merge_cells("A2:H2")
    sub_cell = ws["A2"]
    sub_cell.value = (
        f"Azure Incident Log Pipeline  |  Multi-Agent AI System for Banking Infrastructure  |  "
        f"Generated: {datetime.now().strftime('%d-%b-%Y %H:%M')}"
    )
    sub_cell.font = SUBTITLE_FONT
    sub_cell.alignment = Alignment(horizontal="center", vertical="center")
    ws.row_dimensions[2].height = 22

    # ── Blank row ──
    row = 4

    # ── Header row ──
    for col_idx, header in enumerate(HEADERS, 1):
        cell = ws.cell(row=row, column=col_idx, value=header)
        cell.font = HEADER_FONT
        cell.fill = HEADER_BG
        cell.alignment = center_align
        cell.border = thin_border
    ws.row_dimensions[row].height = 32
    row += 1

    # ── Data rows ──
    serial = 1
    for section_title, features in SECTIONS:
        # Section header row
        ws.merge_cells(start_row=row, start_column=1, end_row=row, end_column=8)
        sec_cell = ws.cell(row=row, column=1, value=section_title)
        sec_cell.font = SECTION_FONT
        sec_cell.fill = SECTION_BG
        sec_cell.alignment = Alignment(horizontal="left", vertical="center")
        sec_cell.border = thin_border
        ws.row_dimensions[row].height = 28
        row += 1

        for idx, (feature, description, stakeholder) in enumerate(features):
            is_alt = idx % 2 == 1
            values = [serial, feature, description, stakeholder, "", "", "", ""]
            for col_idx, val in enumerate(values, 1):
                cell = ws.cell(row=row, column=col_idx, value=val)
                cell.font = BODY_FONT
                cell.alignment = wrap_align if col_idx in (3, 8) else center_align
                cell.border = thin_border
                if is_alt:
                    cell.fill = ALT_ROW

            ws.row_dimensions[row].height = 52
            serial += 1
            row += 1

    # ── Column widths ──
    for col_idx, width in enumerate(COL_WIDTHS, 1):
        ws.column_dimensions[get_column_letter(col_idx)].width = width

    # ── Add data validation for Status column ──
    from openpyxl.worksheet.datavalidation import DataValidation
    status_dv = DataValidation(
        type="list",
        formula1='"Pass,Fail,N/A,Partial"',
        allow_blank=True,
        showDropDown=False,
    )
    status_dv.error = "Please select Pass, Fail, N/A, or Partial"
    status_dv.errorTitle = "Invalid Status"
    status_dv.prompt = "Select validation status"
    status_dv.promptTitle = "Status"
    ws.add_data_validation(status_dv)

    # Apply to all status cells (column E)
    for r in range(5, row):
        cell = ws.cell(row=r, column=5)
        if cell.value == "" or cell.value is None:
            status_dv.add(cell)

    # ── Freeze panes ──
    ws.freeze_panes = "A5"

    # ── Print settings ──
    ws.sheet_properties.pageSetUpPr = openpyxl.worksheet.properties.PageSetupProperties(fitToPage=True)
    ws.page_setup.orientation = "landscape"
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 0

    # ═══════════════════════════════════════════
    #  SUMMARY SHEET
    # ═══════════════════════════════════════════
    ws2 = wb.create_sheet("Summary")
    ws2.merge_cells("A1:D1")
    ws2["A1"].value = "Validation Summary"
    ws2["A1"].font = TITLE_FONT
    ws2["A1"].alignment = Alignment(horizontal="center", vertical="center")
    ws2.row_dimensions[1].height = 36

    summary_headers = ["Stakeholder Domain", "Total Features", "Validated (Pass)", "Pending / Fail"]
    for col_idx, h in enumerate(summary_headers, 1):
        cell = ws2.cell(row=3, column=col_idx, value=h)
        cell.font = HEADER_FONT
        cell.fill = HEADER_BG
        cell.alignment = center_align
        cell.border = thin_border

    summary_data = [
        ("VP – Networking", 7, "", ""),
        ("VP – Cloud", 9, "", ""),
        ("VP – Application", 16, "", ""),
        ("CISO", 10, "", ""),
        ("Cross-Functional", 6, "", ""),
    ]
    for r_idx, (domain, count, validated, pending) in enumerate(summary_data, 4):
        ws2.cell(row=r_idx, column=1, value=domain).font = BODY_FONT
        ws2.cell(row=r_idx, column=2, value=count).font = BODY_FONT
        ws2.cell(row=r_idx, column=3, value=validated).font = BODY_FONT
        ws2.cell(row=r_idx, column=4, value=pending).font = BODY_FONT
        for c in range(1, 5):
            ws2.cell(row=r_idx, column=c).alignment = center_align
            ws2.cell(row=r_idx, column=c).border = thin_border
            if r_idx % 2 == 1:
                ws2.cell(row=r_idx, column=c).fill = ALT_ROW

    # Total row
    total_row = 4 + len(summary_data)
    ws2.cell(row=total_row, column=1, value="TOTAL").font = Font(name="Calibri", size=11, bold=True)
    ws2.cell(row=total_row, column=2, value=f"=SUM(B4:B{total_row-1})").font = Font(name="Calibri", size=11, bold=True)
    for c in range(1, 5):
        ws2.cell(row=total_row, column=c).alignment = center_align
        ws2.cell(row=total_row, column=c).border = thin_border

    # Signoff section
    signoff_row = total_row + 3
    ws2.merge_cells(f"A{signoff_row}:D{signoff_row}")
    ws2[f"A{signoff_row}"].value = "Stakeholder Sign-Off"
    ws2[f"A{signoff_row}"].font = Font(name="Calibri", size=14, bold=True, color=DARK_BLUE)
    ws2[f"A{signoff_row}"].alignment = Alignment(horizontal="center")

    signoff_headers = ["Stakeholder", "Name", "Signature / Approval", "Date"]
    for col_idx, h in enumerate(signoff_headers, 1):
        cell = ws2.cell(row=signoff_row + 1, column=col_idx, value=h)
        cell.font = HEADER_FONT
        cell.fill = HEADER_BG
        cell.alignment = center_align
        cell.border = thin_border

    signoff_roles = [
        "VP – Networking",
        "VP – Cloud",
        "VP – Application",
        "CISO",
    ]
    for r_idx, role in enumerate(signoff_roles, signoff_row + 2):
        ws2.cell(row=r_idx, column=1, value=role).font = BODY_FONT
        for c in range(1, 5):
            ws2.cell(row=r_idx, column=c).alignment = center_align
            ws2.cell(row=r_idx, column=c).border = thin_border
            ws2.row_dimensions[r_idx].height = 30

    for c in range(1, 5):
        ws2.column_dimensions[get_column_letter(c)].width = [30, 25, 30, 18][c - 1]

    # ═══════════════════════════════════════════
    #  SAVE
    # ═══════════════════════════════════════════
    output_path = "CloudGuard_Validation_Checklist.xlsx"
    wb.save(output_path)
    print(f"\n[OK] Validation checklist created: {output_path}")
    print(f"   - {serial - 1} validation items across 5 stakeholder sections")
    print(f"   - Sheets: 'Validation Checklist' + 'Summary'")
    print(f"   - Status dropdown (Pass / Fail / N/A / Partial) enabled")
    print(f"   - Stakeholder sign-off section included\n")


if __name__ == "__main__":
    create_checklist()
